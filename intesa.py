"""
Intesa Sanpaolo import — parses the "Lista Operazioni" account export
(XLSX) into Fluxo preview rows.

The export always covers roughly the last 30 days, so it straddles two
calendar months near a month boundary. To avoid re-importing the tail of
the previous month (already loaded on the prior run), we keep only the
*current month* — the month of the export's "Data fine periodo" (or, as a
fallback, the most recent transaction date).

Classification mirrors revolut.py: deterministic merchant-keyword rules
first, then a fallback that maps Intesa's own "Categoria" column onto
Fluxo's tree. The editable preview handles the tail. Duplicate flagging is
shared with the other bank imports (app._flag_duplicates).
"""
import io
from datetime import datetime, date

# --- Merchant keyword rules (checked first) ---
# (keyword in Operazione/Dettagli, (cat1, cat2, cat3), restrict to tx_type or None)
# First match wins — order matters.
RULES = [
    ('cardmarket',        ('Investments', 'Business', ''), None),
    ('sammelkartenmarkt', ('Investments', 'Business', ''), None),
    ('do invest',         ('Food', 'Snack', ''), None),
    ('ebay',              ('Income', 'Business', ''), 'inflow'),
    ('paypal',            ('Miscellaneous', 'Unknown', ''), None),
    ('american express',  ('Subscriptions', 'Amex', ''), None),
    ('amazon',            ('Miscellaneous', 'Shopping', 'Tech'), None),
    ('uber',              ('Transportation', 'Uber', ''), None),
    ('salvadanaio',       ('Investments', 'Savings', ''), None),
    ('arrotondamento',    ('Investments', 'Roundup', ''), None),
]

# --- Intesa "Categoria" -> Fluxo (cat1, cat2, cat3) fallback map ---
CATEGORY_MAP = {
    'ristoranti e bar':                     ('Food', 'Dinner', ''),
    'generi alimentari e supermercato':     ('Food', 'Groceries', ''),
    'carburanti':                           ('Transportation', 'Gas', ''),
    'trasporti, noleggi, taxi e parcheggi': ('Transportation', '', ''),
    'viaggi e vacanze':                     ('Travel', '', ''),
    'domiciliazioni e utenze':              ('House', 'Charges', ''),
    'spese mediche':                        ('Miscellaneous', 'Personal', 'Health'),
    'tempo libero varie':                   ('Miscellaneous', 'Entertainment', ''),
    'investimenti, bdr e salvadanaio':      ('Investments', 'Savings', ''),
    'disinvestimenti, bdr e salvadanaio':   ('Investments', 'Savings', ''),
    'bonifici ricevuti':                    ('Income', '', ''),
    'entrate varie':                        ('Income', '', ''),
    'bonifici in uscita':                   ('Miscellaneous', 'Gifts', ''),
    'imposte, bolli e commissioni':         ('Miscellaneous', 'Unknown', ''),
    'addebiti vari':                        ('Miscellaneous', 'Unknown', ''),
    'altre uscite':                         ('Miscellaneous', 'Unknown', ''),
    'associazioni':                         ('Miscellaneous', 'Unknown', ''),
    'addebiti nexi e carte non del gruppo intesa sanpaolo':
                                            ('Miscellaneous', 'Unknown', ''),
}

DEFAULT_OUTFLOW = ('Miscellaneous', 'Unknown', '')
DEFAULT_INFLOW = ('Income', '', '')

# Generic "Operazione" labels that carry no merchant — for card payments the
# real merchant sits in the "Dettagli" column, so prefer that. (Bonifici and
# addebiti keep the Operazione text, which names the counterparty.)
GENERIC_OPS = ('pagamento',)


def _merchant(op, det):
    """Pick the most descriptive label for the preview."""
    if op and not any(op.lower().startswith(p) for p in GENERIC_OPS):
        return op
    return det or op

# Movements between own accounts — not real spend; pre-deselected in preview.
# (Top-ups to the linked Revolut account, which are imported separately.)
INTERNAL_PATTERNS = ('revolut',)


def classify(text, intesa_category, tx_type):
    t = (text or '').lower()
    for key, path, only in RULES:
        if key in t and (only is None or only == tx_type):
            return path
    mapped = CATEGORY_MAP.get((intesa_category or '').strip().lower())
    if mapped:
        return mapped
    return DEFAULT_INFLOW if tx_type == 'inflow' else DEFAULT_OUTFLOW


def _parse_it_date(val):
    """Parse a dd/mm/yyyy string or a datetime cell into a date, else None."""
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if not val:
        return None
    try:
        return datetime.strptime(str(val).strip(), '%d/%m/%Y').date()
    except ValueError:
        return None


def _read_rows(file_bytes, filename):
    """Return a list of row-lists from an Intesa Sanpaolo XLSX export."""
    name = (filename or '').lower()
    if not name.endswith(('.xlsx', '.xls')):
        raise ValueError('Not an Intesa Sanpaolo export — expected an Excel (.xlsx) file')
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _find_header(rows):
    """Locate the transactions header row and return (row_index, column map)."""
    for i, r in enumerate(rows):
        cells = [str(c).strip().lower() if c is not None else '' for c in r]
        if 'data' in cells and 'importo' in cells and 'operazione' in cells:
            col = {name: cells.index(name) for name in
                   ('data', 'operazione', 'dettagli', 'contabilizzazione',
                    'categoria', 'importo') if name in cells}
            return i, col
    raise ValueError('Not an Intesa Sanpaolo export — transaction header not found')


def _reference_month(rows, tx_dates):
    """The month to keep: the export's end-of-period, else the latest tx date."""
    for r in rows:
        for j, c in enumerate(r):
            if c is not None and 'data fine periodo' in str(c).strip().lower():
                for nxt in r[j + 1:]:
                    d = _parse_it_date(nxt)
                    if d:
                        return d.year, d.month
    if tx_dates:
        d = max(tx_dates)
        return d.year, d.month
    return None


def parse_export(file_bytes, filename):
    """Parse a raw Intesa Sanpaolo export into preview rows (same dict shape
    revolut.parse_export produces, plus method). Only current-month rows are
    returned; pending (not-yet-booked) rows are pre-deselected."""
    rows = _read_rows(file_bytes, filename)
    if not rows:
        raise ValueError('Empty file — is it an Intesa Sanpaolo export?')

    header_i, col = _find_header(rows)
    i_date = col.get('data')
    i_op = col.get('operazione')
    i_det = col.get('dettagli')
    i_cont = col.get('contabilizzazione')
    i_cat = col.get('categoria')
    i_amt = col.get('importo')
    if i_date is None or i_amt is None:
        raise ValueError('Not an Intesa Sanpaolo export — missing Data/Importo columns')

    body = rows[header_i + 1:]

    def g(r, i):
        return r[i] if i is not None and i < len(r) and r[i] is not None else ''

    # First pass: collect dates to establish the current month.
    parsed = []
    for r in body:
        d = _parse_it_date(g(r, i_date))
        if d is None:
            continue
        try:
            amount = float(g(r, i_amt))
        except (ValueError, TypeError):
            continue
        parsed.append((r, d, amount))

    ref = _reference_month(rows, [d for _, d, _ in parsed])

    out = []
    for r, d, amount in parsed:
        # Keep only the current month — the export's tail of the previous
        # month was already imported on the prior run.
        if ref and (d.year, d.month) != ref:
            continue

        op = str(g(r, i_op)).strip()
        det = str(g(r, i_det)).strip()
        merchant = _merchant(op, det)
        intesa_cat = str(g(r, i_cat)).strip()
        booked = str(g(r, i_cont)).strip().upper()  # 'SI' booked, 'NO' pending
        tx_type = 'inflow' if amount > 0 else 'outflow'

        include, note = True, ''
        if booked == 'NO':
            include, note = False, 'pending (not yet booked)'
        elif any(p in merchant.lower() for p in INTERNAL_PATTERNS):
            include, note = False, 'internal transfer'

        c1, c2, c3 = classify(f'{op} {det}', intesa_cat, tx_type)
        out.append({
            'date': d.isoformat(),
            'merchant': merchant,
            'transaction_type': tx_type,
            'amount': round(abs(amount), 2),
            'category_1': c1, 'category_2': c2, 'category_3': c3,
            'method': 'Intesa',
            'include': include,
            'note': note,
        })

    if not out:
        raise ValueError('No current-month transactions found in the export')
    return out
